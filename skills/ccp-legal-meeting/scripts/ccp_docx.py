# -*- coding: utf-8 -*-
"""党支部会议资料 docx/pptx helper（可复用机械部分）。

用法：import 本模块后调用 new_doc/para/title/sign_sheet/retext_clone/add_footer_pagenum。
内容（议程、发言、日期）由调用方按当月学习资料定制。
依赖：python-docx、python-pptx。
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FT_TITLE = '方正小标宋简体'
FT_HEI = '黑体'
FT_FANG = '仿宋_GB2312'
FT_ASCII = 'Times New Roman'


def set_run(run, ea=FT_FANG, size=15, bold=False):
    run.font.name = FT_ASCII
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.append(rf)
    rf.set(qn('w:eastAsia'), ea)
    run.font.size = Pt(size)
    run.font.bold = bold


def para(doc, text='', ea=FT_FANG, size=15, bold=False, align='justify', indent=True,
         line=22, before=0, after=0):
    """标准段落。indent=True 时首行缩进2字符（get_or_add_ind，勿手动 append）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pPr = p._element.get_or_add_pPr()
        e = pPr.get_or_add_ind()
        e.set(qn('w:firstLineChars'), '200')
        e.set(qn('w:firstLine'), '600')
    if text:
        run = p.add_run(text)
        set_run(run, ea, size, bold)
    return p


def title(doc, lines):
    """文档标题：方正小标宋 22pt 加粗居中，固定行距28.9磅，可传多行。"""
    for t in lines:
        para(doc, t, ea=FT_TITLE, size=22, bold=True, align='center', indent=False, line=28.9)


def add_footer_pagenum(doc):
    """页脚居中 PAGE 域（≥2页 docx 必加）。"""
    footer = doc.sections[0].footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run()
    fld = OxmlElement('w:fldChar'); fld.set(qn('w:fldCharType'), 'begin'); r1._element.append(fld)
    r2 = p.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' PAGE '
    r2._element.append(it)
    r3 = p.add_run()
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end'); r3._element.append(fld2)
    for r in p.runs:
        r.font.name = FT_ASCII
        r.font.size = Pt(10.5)


def new_doc():
    """A4、边距上下2.54/左右3.17cm、Normal=仿宋_GB2312+Times New Roman。"""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.17)
    st = doc.styles['Normal']
    st.font.name = FT_ASCII
    st._element.get_or_add_rPr()
    rf = st._element.rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); st._element.rPr.append(rf)
    rf.set(qn('w:eastAsia'), FT_FANG)
    return doc


def save(doc, path, pagenum=False):
    if pagenum:
        add_footer_pagenum(doc)
    doc.save(path)
    print('生成:', path)


def notice(doc_title_lines, head, date_str, participants, items, rules,
           topic_line=None, closing=('党支部', None), notice_date=''):
    """通知骨架。items=[议程字符串...]；rules=[纪律要求...]；topic_line=党课主题行(可选)。"""
    title(doc, doc_title_lines)
    para(doc, '')
    para(doc, head)
    para(doc, '一、会议时间', ea=FT_HEI, bold=True)
    para(doc, date_str + '。')
    para(doc, '二、会议地点', ea=FT_HEI, bold=True)
    para(doc, '三楼会议室。')
    para(doc, '三、参会人员', ea=FT_HEI, bold=True)
    para(doc, participants + '。')
    para(doc, '四、会议内容', ea=FT_HEI, bold=True)
    if topic_line:
        para(doc, topic_line)
    for i, it in enumerate(items, 1):
        para(doc, f'{i}、{it}；')
    para(doc, '五、纪律要求', ea=FT_HEI, bold=True)
    for r in rules:
        para(doc, r)
    para(doc, '')
    para(doc, closing[0], align='right')
    para(doc, notice_date, align='right')
    return doc


def record_base(doc, docname_lines, name, date_week, attendees_lines, body_size=15):
    """会议记录头部：标题+信息行。"""
    title(doc, docname_lines)
    para(doc, '')
    for t in ['会议名称：' + name, '开会时间：' + date_week, '会议地点：三楼会议室']:
        para(doc, t, indent=False, size=body_size)
    for t in attendees_lines:
        para(doc, t, indent=False, size=body_size)
    para(doc, '主 持 人：×××                记录人：×××', indent=False, size=body_size)
    para(doc, '')
    return doc


def sec(doc, text, size=15):
    """记录内议程标题（黑体加粗）。"""
    para(doc, text, ea=FT_HEI, bold=True, size=size)


def sign_sheet(path, subject, people, party_set, date_str, with_status=False):
    """签到表：title + (时间/地点/主题/主持人/参加人员) + 姓名|签到双列网格。
    people=按展示顺序的名单；party_set=党员集合（判定政治面貌）。"""
    doc = new_doc()
    para(doc, '××××有限公司签到表', ea=FT_TITLE, size=18, bold=True,
         align='center', indent=False, line=None)
    para(doc, '')
    names = list(people)
    half = (len(names) + 1) // 2
    left = names[:half]
    ncols = 6 if with_status else 4
    table = doc.add_table(rows=half + 5, cols=ncols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    def cell(r, c, text, bold=False, ea=FT_FANG, size=12):
        p = table.rows[r].cells[c].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(text), ea, size, bold)

    def merged(r, c1, c2, text, align_left=False):
        m = table.rows[r].cells[c1].merge(table.rows[r].cells[c2])
        p = m.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if align_left else WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(text), FT_FANG, 12)

    if with_status:
        cell(0, 0, '时  间', bold=True, ea=FT_HEI); merged(0, 1, 2, date_str)
        cell(0, 3, '地  点', bold=True, ea=FT_HEI); merged(0, 4, 5, '三楼会议室')
        cell(1, 0, '主  题', bold=True, ea=FT_HEI); merged(1, 1, 5, subject)
        cell(2, 0, '主持人', bold=True, ea=FT_HEI); merged(2, 1, 5, '×××')
        merged(3, 0, 5, '参加人员：', align_left=True)
        for c, h in enumerate(['姓名', '政治面貌', '签到', '姓名', '政治面貌', '签到']):
            cell(4, c, h, bold=True, ea=FT_HEI)
        start = 5
        for i in range(half):
            rr = start + i
            cell(rr, 0, left[i])
            cell(rr, 1, '中共党员' if left[i] in party_set else '群众')
            j = half + i
            if j < len(names):
                cell(rr, 3, names[j])
                cell(rr, 4, '中共党员' if names[j] in party_set else '群众')
    else:
        cell(0, 0, '时  间', bold=True, ea=FT_HEI); cell(0, 1, date_str)
        cell(0, 2, '地  点', bold=True, ea=FT_HEI); cell(0, 3, '三楼会议室')
        cell(1, 0, '主  题', bold=True, ea=FT_HEI); merged(1, 1, 3, subject)
        cell(2, 0, '主持人', bold=True, ea=FT_HEI); merged(2, 1, 3, '×××')
        merged(3, 0, 3, '参加人员：', align_left=True)
        for c, h in enumerate(['姓名', '签到', '姓名', '签到']):
            cell(4, c, h, bold=True, ea=FT_HEI)
        start = 5
        for i in range(half):
            rr = start + i
            cell(rr, 0, left[i])
            j = half + i
            if j < len(names):
                cell(rr, 2, names[j])
    doc.save(path)
    print('生成:', path)


def replace_para(p, new_text):
    """模板段落替换文字（保第一个 run 格式）。"""
    runs = p.runs
    if not runs:
        p.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ''


def cell_add(cl, text):
    """表格单元格填值（已非空则跳过，防重复叠加）。"""
    p = cl.paragraphs[0]
    if not p.text.strip():
        p.add_run(text)


# ---------------- pptx ----------------

def _norm(s):
    return ''.join(s.split())


def retext_clone(tf, lines):
    """文本框保格式替换为多行；行数多于段落时 deepcopy 克隆最后有文字的段落。"""
    import copy
    paras = tf.paragraphs
    src_para = None
    for para in paras:
        if para.text.strip():
            src_para = para
    while len(tf.paragraphs) < len(lines) and src_para is not None:
        new_el = copy.deepcopy(src_para._element)
        tf.paragraphs[-1]._element.addnext(new_el)
    for i, para in enumerate(tf.paragraphs):
        if i < len(lines):
            runs = para.runs
            if not runs:
                para.text = lines[i]
            else:
                runs[0].text = lines[i]
                for r in runs[1:]:
                    r.text = ''
        else:
            for r in para.runs:
                r.text = ''


def apply_ppt_mapping(path, mapping):
    """按 {旧文本: [新行...]} 替换 pptx 文字（去空白归一化匹配，保字体背景图）。"""
    from pptx import Presentation
    pr = Presentation(path)
    done = set()
    for slide in pr.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            key = _norm(sh.text_frame.text)
            for old, new in mapping.items():
                if _norm(old) == key and old not in done:
                    retext_clone(sh.text_frame, new)
                    done.add(old)
    pr.save(path)
    for old in mapping:
        print(('OK  ' if old in done else 'MISS'), old.replace('\n', '/')[:40])
    return done
