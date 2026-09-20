# -*- coding: utf-8 -*-
"""
中文公文 / 诉讼文书 Word 排版示例（python-docx）
==============================================

演示如何按本仓库规范生成一份合规的 .docx：

  · A4 页面，上下 2.54cm、左右 3.17cm 边距
  · 西文与数字统一 Times New Roman
  · 中文按层级：题目 方正小标宋简体 / 一级 黑体 / 二级 楷体 / 三级与正文 仿宋_GB2312
  · 正文：首行缩进 2 字符、固定行距 22 磅、两端对齐
  · 标题使用 Word 内置 heading 样式（导航窗格可见）
  · 页脚居中页码（PAGE 域）

用法：
    pip install python-docx
    python example_generate.py [输出文件名.docx]

要点提示（对应 SKILL.md 第五节）：删除 theme 字体属性、关闭内置标题的斜体、
段落显式写 w:ind、w:ind 必须排在 w:jc 之前。
"""
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- 常量
FONT_CN_TITLE = "方正小标宋简体"   # 题目
FONT_CN_H1 = "黑体"                # 一级标题（一、）
FONT_CN_H2 = "楷体"                # 二级标题（（一））
FONT_CN_BODY = "仿宋_GB2312"       # 三级标题与正文
FONT_EN = "Times New Roman"        # 西文与数字（全局）

LINE_EXACT_22PT = Pt(22)           # 正文固定行距
FIRST_LINE_INDENT = Pt(30)         # ≈2 字符（15pt × 2）


# ---------------------------------------------------------------- 工具函数
def set_font(run, cn_font, size_pt, bold=False, italic=False):
    """给 run 设置中英文字体、字号、加粗，并清除主题字体残留。"""
    run.font.name = FONT_EN                        # ascii / hAnsi
    run.font.size = Pt(size_pt)
    run.font.bold = bold

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), cn_font)

    # 关键：删除 theme 字体属性，否则 Word 可能回退成宋体 / MS Gothic
    for attr in list(rfonts.attrib):
        if attr.endswith("Theme"):
            del rfonts.attrib[attr]

    # 关键：显式关闭斜体（内置 Heading 4 等自带 w:i）
    for tag in ("w:i", "w:iCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), "1" if italic else "0")

    # 全文黑色，避免继承内置标题样式的蓝色
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")


def set_para_format(para, line_exact=True, first_line_chars=2,
                    align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0):
    """设置段落：固定行距、首行缩进（按字符数转 twips）、对齐、段前段后。"""
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line_exact:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = LINE_EXACT_22PT
    para.alignment = align

    ppr = para._p.get_or_add_pPr()
    if first_line_chars:
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            # 关键：w:ind 必须排在 w:jc 之前，否则 XSD 报 "ind not expected"
            jc = ppr.find(qn("w:jc"))
            if jc is not None:
                jc.addprevious(ind)
            else:
                ppr.append(ind)
        # 15pt 字号下 1 字符 = 300 twips
        ind.set(qn("w:firstLine"), str(int(first_line_chars * 300)))
        ind.set(qn("w:firstLineChars"), str(int(first_line_chars * 100)))
    else:
        # 顶格：显式写 0，避免继承 docDefaults
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            jc = ppr.find(qn("w:jc"))
            if jc is not None:
                jc.addprevious(ind)
            else:
                ppr.append(ind)
        ind.set(qn("w:firstLine"), "0")
        ind.set(qn("w:firstLineChars"), "0")


def add_footer_page_number(section):
    """页脚居中 PAGE 域（Word「页面底端居中」）。"""
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    run = para.add_run()
    set_font(run, FONT_CN_BODY, 14)

    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(instr); run._r.append(end)


def force_style_font(doc, style_name, cn_font, size_pt=None, bold=None):
    """把样式自身的字体写死（并清除 theme 属性），作为兜底。"""
    try:
        style = doc.styles[style_name]
    except KeyError:
        return
    style.font.name = FONT_EN
    if size_pt:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), cn_font)
    for attr in list(rfonts.attrib):
        if attr.endswith("Theme"):
            del rfonts.attrib[attr]


# ---------------------------------------------------------------- 主流程
def build(path="示例-工作证明.docx"):
    doc = Document()

    # 页面
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)

    # 样式兜底（防止 Word 用 theme 字体渲染）
    force_style_font(doc, "Normal", FONT_CN_BODY, 15)
    force_style_font(doc, "Heading 1", FONT_CN_TITLE, 22, True)
    force_style_font(doc, "Heading 2", FONT_CN_H1, 15, True)
    force_style_font(doc, "Heading 3", FONT_CN_H2, 15, True)
    force_style_font(doc, "Heading 4", FONT_CN_BODY, 15, True)

    # 题目（Heading 1 → 导航可见）
    p = doc.add_heading(level=1)
    r = p.add_run("工 作 证 明")
    set_font(r, FONT_CN_TITLE, 22, bold=True)
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_chars=0, after=15.6)

    # 一级标题
    p = doc.add_heading(level=2)
    r = p.add_run("一、基本情况")
    set_font(r, FONT_CN_H1, 15, bold=True)
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=7.8, after=7.8)

    # 正文
    body_texts = [
        "兹证明×××同志自××××年××月至今在我单位工作，现任××岗位。",
        "该同志在职期间工作认真负责，能够独立完成交办的各项工作任务。",
    ]
    for text in body_texts:
        p = doc.add_paragraph()
        r = p.add_run(text)
        set_font(r, FONT_CN_BODY, 15)
        set_para_format(p, before=7.8, after=7.8)

    # 二级标题
    p = doc.add_heading(level=3)
    r = p.add_run("（一）主要工作内容")
    set_font(r, FONT_CN_H2, 15, bold=True)
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT)

    # 三级标题
    p = doc.add_heading(level=4)
    r = p.add_run("1. 日常事务处理")
    set_font(r, FONT_CN_BODY, 15, bold=True)
    set_para_format(p, align=WD_ALIGN_PARAGRAPH.LEFT)

    p = doc.add_paragraph()
    r = p.add_run("包括文件收发、会议记录、资料归档等常规工作。")
    set_font(r, FONT_CN_BODY, 15)
    set_para_format(p)

    # 落款（右对齐）
    for text in ("", "证明单位（盖章）：××××××", "××××年××月××日"):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text)
            set_font(r, FONT_CN_BODY, 15)
        set_para_format(p, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # 页脚页码
    add_footer_page_number(sec)

    doc.save(path)
    print("已生成:", path)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "示例-工作证明.docx")
